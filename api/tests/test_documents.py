"""Faz 2: yükleme doğrulama, metin çıkarma, kullanıcı yalıtımı, limitler."""
import io
import time
import uuid
import zipfile
from datetime import datetime

import jwt
import pytest
from docx import Document
from fastapi.testclient import TestClient
from pptx import Presentation

from app import extract as ex
from app.config import get_settings
from app.main import app
from app.store import DuplicateError, get_store

SECRET = "x" * 40


class FakeStore:
    def __init__(self):
        self.docs: dict[str, dict] = {}
        self.usage: list[tuple[str, str]] = []

    async def find_by_hash(self, user_id, h):
        for d in self.docs.values():
            if d["user_id"] == user_id and d["content_hash"] == h:
                return self._pub(d)

    async def create_document(self, row):
        if await self.find_by_hash(row["user_id"], row["content_hash"]):
            raise DuplicateError()
        d = {"id": str(uuid.uuid4()), "created_at": datetime.now().isoformat(), "raw_text": None, "error": None,
             "page_count": None, "char_count": None, "used_ocr": False, "truncated": False, "note": None, **row}
        self.docs[d["id"]] = d
        return dict(d)

    async def update_document(self, doc_id, user_id, values):
        d = self.docs.get(doc_id)
        if d and d["user_id"] == user_id:
            d.update(values)

    async def get_document(self, doc_id, user_id):
        d = self.docs.get(doc_id)
        return dict(d) if d and d["user_id"] == user_id else None

    async def list_documents(self, user_id):
        return [self._pub(d) for d in self.docs.values() if d["user_id"] == user_id]

    async def delete_document(self, doc_id, user_id):
        d = self.docs.get(doc_id)
        if d and d["user_id"] == user_id:
            del self.docs[doc_id]
            return True
        return False

    async def count_uploads_since(self, user_id, since):
        return sum(1 for u, a in self.usage if u == user_id and a == "upload")

    async def log_usage(self, user_id, action, units=0):
        self.usage.append((user_id, action))

    @staticmethod
    def _pub(d):
        return {k: v for k, v in d.items() if k != "raw_text"}


@pytest.fixture
def store(monkeypatch):
    monkeypatch.setenv("SUPABASE_JWT_SECRET", SECRET)
    get_settings.cache_clear()
    fs = FakeStore()
    app.dependency_overrides[get_store] = lambda: fs
    yield fs
    app.dependency_overrides.clear()
    get_settings.cache_clear()


client = TestClient(app)


def tok(sub="u1", role=None):
    c = {"sub": sub, "email": f"{sub}@x.com", "aud": "authenticated", "exp": int(time.time()) + 3600}
    if role:
        c["app_metadata"] = {"role": role}
    return {"Authorization": "Bearer " + jwt.encode(c, SECRET, algorithm="HS256")}


# ---------- örnek dosya üreticileri ----------
def make_docx(text="Hucre zari secici gecirgendir ve madde alisverisini duzenler."):
    d = Document()
    d.add_heading("Biyoloji Notlari", 1)
    d.add_paragraph(text)
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text, t.cell(1, 0).text, t.cell(1, 1).text = "Tasima", "Enerji", "Pasif", "Yok"
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()


def make_pptx():
    p = Presentation()
    s = p.slides.add_slide(p.slide_layouts[1])
    s.shapes.title.text = "Fotosentez"
    s.placeholders[1].text = "Isik enerjisi kimyasal enerjiye donusur."
    s.notes_slide.notes_text_frame.text = "Kloroplast vurgulanacak"
    b = io.BytesIO()
    p.save(b)
    return b.getvalue()


def make_pdf(text="Mitokondri hucrenin enerji santralidir ve ATP uretir."):
    content = f"BT /F1 18 Tf 72 700 Td ({text}) Tj ET" if text else ""
    objs = ["<< /Type /Catalog /Pages 2 0 R >>", "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
            f"<< /Length {len(content)} >>\nstream\n{content}\nendstream", "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    out, offs = b"%PDF-1.4\n", []
    for i, o in enumerate(objs, 1):
        offs.append(len(out))
        out += f"{i} 0 obj\n{o}\nendobj\n".encode()
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    for off in offs:
        out += f"{off:010d} 00000 n \n".encode()
    return out + f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()


def up(name, data, headers=None):
    return client.post("/api/documents", files={"file": (name, data)}, headers=headers or tok())


def text_of(doc_id, headers=None):
    return client.get(f"/api/documents/{doc_id}", headers=headers or tok()).json()


# ---------- erişim ----------
@pytest.mark.parametrize("method,path", [("get", "/api/documents"), ("get", f"/api/documents/{uuid.uuid4()}"),
                                         ("delete", f"/api/documents/{uuid.uuid4()}")])
def test_login_required(store, method, path):
    assert getattr(client, method)(path).status_code == 401


def test_upload_requires_login(store):
    assert client.post("/api/documents", files={"file": ("a.pdf", make_pdf())}).status_code == 401


# ---------- çıkarma ----------
def test_docx_flow(store):
    r = up("notlar.docx", make_docx())
    assert r.status_code == 202
    d = text_of(r.json()["id"])
    assert d["status"] == "ready" and d["kind"] == "docx"
    assert "## Biyoloji Notlari" in d["text_preview"] and "Tasima | Enerji" in d["text_preview"]
    assert "raw_text" not in d


def test_pptx_flow(store):
    d = text_of(up("sunum.pptx", make_pptx()).json()["id"])
    assert d["status"] == "ready" and d["page_count"] == 1
    assert "Fotosentez" in d["text_preview"] and "Konuşmacı notu: Kloroplast" in d["text_preview"]


def test_pdf_flow(store):
    d = text_of(up("ders.pdf", make_pdf()).json()["id"])
    assert d["status"] == "ready" and d["page_count"] == 1 and "Mitokondri" in d["text_preview"]
    assert d["used_ocr"] is False


def test_scanned_pdf_uses_ocr(store, monkeypatch):
    monkeypatch.setattr(ex, "get_ocr_fn", lambda: (lambda img: "OCR ile okunan uzun bir metin parcasi"))
    d = text_of(up("tarama.pdf", make_pdf(None)).json()["id"])
    assert d["status"] == "ready" and d["used_ocr"] is True and "OCR ile okunan" in d["text_preview"]
    assert "OCR" in d["note"]


def test_scanned_pdf_without_ocr_fails_cleanly(store, monkeypatch):
    monkeypatch.setattr(ex, "get_ocr_fn", lambda: None)
    d = text_of(up("tarama.pdf", make_pdf(None)).json()["id"])
    assert d["status"] == "failed" and "taranmış" in d["error"]


def test_real_tesseract_if_available():
    if not ex.ocr_available():
        pytest.skip("tesseract yok")
    from PIL import Image, ImageDraw, ImageFont
    img = Image.new("RGB", (900, 200), "white")
    try:
        font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 56)
    except OSError:
        pytest.skip("yazı tipi yok")
    ImageDraw.Draw(img).text((20, 50), "Cell membrane test", fill="black", font=font)
    assert "membrane" in ex.tesseract_ocr(img).lower()


# ---------- doğrulama ----------
def test_rejects_wrong_extension(store):
    assert up("not.txt", b"merhaba dunya").status_code == 415


def test_rejects_fake_pdf(store):
    assert up("sahte.pdf", b"bu bir pdf degil" * 50).status_code == 415


def test_rejects_extension_mismatch(store):
    assert up("yanlis.docx", make_pdf()).status_code == 415


def test_rejects_empty(store):
    assert up("bos.pdf", b"").status_code == 422


def test_rejects_too_large(store, monkeypatch):
    monkeypatch.setenv("MAX_UPLOAD_MB", "1")
    get_settings.cache_clear()
    assert up("buyuk.pdf", b"%PDF-" + b"0" * (2 * 1024 * 1024)).status_code == 413


def test_corrupt_pdf_fails_cleanly(store):
    d = text_of(up("bozuk.pdf", b"%PDF-1.4 bozuk icerik").json()["id"])
    assert d["status"] == "failed" and d["error"]


def test_zip_bomb_rejected(store, monkeypatch):
    monkeypatch.setenv("MAX_UNCOMPRESSED_MB", "1")
    get_settings.cache_clear()
    b = io.BytesIO()
    with zipfile.ZipFile(b, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", "<Types/>")
        z.writestr("word/document.xml", "0" * (20 * 1024 * 1024))
    d = text_of(up("bomba.docx", b.getvalue()).json()["id"])
    assert d["status"] == "failed" and "güvenlik" in d["error"]


def test_path_in_filename_is_stripped(store):
    r = up("../../etc/passwd.docx", make_docx())
    assert r.status_code == 202 and r.json()["title"] == "passwd.docx"


# ---------- önbellek, limit, yalıtım ----------
def test_same_file_is_cached(store):
    a = up("a.docx", make_docx()).json()
    b = up("kopya.docx", make_docx())
    assert b.status_code == 200 and b.json()["cached"] is True and b.json()["id"] == a["id"]
    assert len(store.docs) == 1 and len(store.usage) == 1


def test_daily_limit(store, monkeypatch):
    monkeypatch.setenv("DAILY_LIMIT_PER_USER", "2")
    get_settings.cache_clear()
    assert up("1.docx", make_docx("Birinci belge icerigi burada yer alir.")).status_code == 202
    assert up("2.docx", make_docx("Ikinci belge icerigi burada yer alir.")).status_code == 202
    assert up("3.docx", make_docx("Ucuncu belge icerigi burada yer alir.")).status_code == 429


def test_users_are_isolated(store):
    a = up("a.docx", make_docx(), tok("alice")).json()["id"]
    assert client.get(f"/api/documents/{a}", headers=tok("bob")).status_code == 404
    assert client.delete(f"/api/documents/{a}", headers=tok("bob")).status_code == 404
    assert client.get("/api/documents", headers=tok("bob")).json() == []
    assert len(client.get("/api/documents", headers=tok("alice")).json()) == 1
    assert client.get(f"/api/documents/{a}", headers=tok("alice")).status_code == 200


def test_delete_and_reupload(store):
    a = up("a.docx", make_docx()).json()["id"]
    assert client.delete(f"/api/documents/{a}", headers=tok()).status_code == 204
    assert client.get(f"/api/documents/{a}", headers=tok()).status_code == 404
    assert up("a.docx", make_docx()).status_code == 202


def test_failed_doc_is_retried_on_reupload(store, monkeypatch):
    monkeypatch.setattr(ex, "get_ocr_fn", lambda: None)
    pdf = make_pdf(None)
    first = up("t.pdf", pdf).json()["id"]
    assert text_of(first)["status"] == "failed"
    second = up("t.pdf", pdf)
    assert second.status_code == 202 and second.json()["id"] != first
