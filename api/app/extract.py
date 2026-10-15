"""Belgeden metin çıkarma. Dosyalar diske yazılmaz, bellekte işlenir.
Tüm ExtractError mesajları kullanıcıya gösterilebilir (Türkçe) metinlerdir."""
import io
import re
import shutil
import zipfile
from dataclasses import dataclass
from functools import lru_cache

from .config import get_settings

MAGIC_PDF = b"%PDF-"
MAGIC_ZIP = b"PK\x03\x04"
KINDS = {"pdf", "docx", "pptx"}
MAX_ZIP_ENTRIES = 3000
MIN_PAGE_CHARS = 25  # bundan az metin içeren PDF sayfası görsel/taranmış sayılır


class ExtractError(Exception):
    pass


class UnsupportedFile(ExtractError):
    pass


@dataclass
class Extracted:
    text: str
    page_count: int | None = None
    used_ocr: bool = False
    truncated: bool = False
    note: str | None = None


def clean_filename(name: str | None) -> str:
    name = (name or "belge").replace("\\", "/").split("/")[-1]
    name = re.sub(r"[\x00-\x1f]", "", name).strip() or "belge"
    return name[:200]


def detect_kind(filename: str, data: bytes) -> str:
    """Uzantıya VE dosyanın gerçek başlığına bakar (uzantı sahteciliğine karşı)."""
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext not in KINDS:
        raise UnsupportedFile("Desteklenen türler: PDF, Word (.docx) ve PowerPoint (.pptx).")
    ok = data[:1024].lstrip().startswith(MAGIC_PDF) if ext == "pdf" else data.startswith(MAGIC_ZIP)
    if not ok:
        raise UnsupportedFile("Dosyanın içeriği uzantısıyla uyuşmuyor veya dosya geçerli değil.")
    return ext


# ---------------- OCR ----------------
def ocr_available() -> bool:
    return shutil.which("tesseract") is not None


@lru_cache
def _ocr_lang() -> str:
    try:
        import pytesseract
        langs = set(pytesseract.get_languages())
    except Exception:
        return "eng"
    want = [l for l in ("tur", "eng") if l in langs]
    return "+".join(want) or "eng"


def tesseract_ocr(img) -> str:
    import pytesseract
    return pytesseract.image_to_string(img, lang=_ocr_lang(), timeout=90)


def get_ocr_fn():
    """Testlerde değiştirilebilsin diye çağrı anında çözülür."""
    return tesseract_ocr if ocr_available() else None


# ---------------- güvenlik ----------------
def _check_zip(data: bytes, required: str, what: str) -> None:
    s = get_settings()
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            infos = z.infolist()
            total = sum(i.file_size for i in infos)
            if len(infos) > MAX_ZIP_ENTRIES or total > s.max_uncompressed_mb * 1024 * 1024:
                raise ExtractError("Dosya güvenlik sınırlarını aşıyor (açıldığında aşırı büyük).")
            names = set(z.namelist())
    except zipfile.BadZipFile:
        raise ExtractError(f"Dosya bozuk veya geçerli bir {what} belgesi değil.")
    if "[Content_Types].xml" not in names or required not in names:
        raise ExtractError(f"Bu dosya geçerli bir {what} belgesi değil.")


# ---------------- biçimlere göre çıkarma ----------------
def _pdf(data: bytes) -> Extracted:
    import pypdfium2 as pdfium

    s = get_settings()
    try:
        pdf = pdfium.PdfDocument(data)
    except Exception as e:
        if "password" in str(e).lower():
            raise ExtractError("Parola korumalı PDF'ler desteklenmiyor.")
        raise ExtractError("PDF dosyası açılamadı (bozuk olabilir).")
    try:
        n = len(pdf)
        if n > s.max_pages:
            raise ExtractError(f"PDF en fazla {s.max_pages} sayfa olabilir (bu dosya {n} sayfa).")
        ocr = get_ocr_fn()
        parts: list[str] = []
        ocr_pages = image_pages = 0
        for i in range(n):
            page = pdf[i]
            try:
                tp = page.get_textpage()
                try:
                    t = tp.get_text_range() or ""
                finally:
                    tp.close()
                if len(t.strip()) < MIN_PAGE_CHARS:
                    image_pages += 1
                    if ocr is not None and ocr_pages < s.ocr_max_pages:
                        try:
                            t = ocr(page.render(scale=200 / 72).to_pil().convert("RGB")) or ""
                            ocr_pages += 1
                        except Exception:
                            t = ""
                parts.append(t)
            finally:
                page.close()
        note = None
        if ocr_pages:
            note = f"{ocr_pages} görsel sayfa OCR ile okundu; okuma hataları olabilir."
            if image_pages > ocr_pages:
                note += f" OCR en fazla {s.ocr_max_pages} sayfa için uygulanır, kalan görsel sayfalar atlandı."
        return Extracted("\n\n".join(parts), page_count=n, used_ocr=ocr_pages > 0, note=note)
    finally:
        pdf.close()


def _dedupe_cells(cells: list[str]) -> list[str]:
    out: list[str] = []
    for c in cells:
        if not out or c != out[-1]:  # birleştirilmiş hücreler tekrar etmesin
            out.append(c)
    return out


def _docx(data: bytes) -> Extracted:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    _check_zip(data, "word/document.xml", "Word")
    try:
        doc = Document(io.BytesIO(data))
    except Exception:
        raise ExtractError("Word dosyası açılamadı (bozuk olabilir).")
    out: list[str] = []
    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(child, doc)
            t = p.text.strip()
            if not t:
                continue
            style = ((p.style.name if p.style is not None else "") or "").lower()
            m = re.match(r"(heading|başlık)\s*(\d)", style)
            if style == "title":
                t = "# " + t
            elif m:
                t = "#" * min(int(m.group(2)) + 1, 6) + " " + t
            out.append(t)
        elif tag == "tbl":
            for row in Table(child, doc).rows:
                cells = _dedupe_cells([c.text.strip().replace("\n", " ") for c in row.cells])
                if any(cells):
                    out.append(" | ".join(cells))
    return Extracted("\n".join(out))


def _walk_shapes(shapes):
    for sh in shapes:
        if getattr(sh, "shape_type", None) == 6:  # grup
            yield from _walk_shapes(sh.shapes)
        else:
            yield sh


def _pptx(data: bytes) -> Extracted:
    from pptx import Presentation

    _check_zip(data, "ppt/presentation.xml", "PowerPoint")
    try:
        prs = Presentation(io.BytesIO(data))
    except Exception:
        raise ExtractError("PowerPoint dosyası açılamadı (bozuk olabilir).")
    out: list[str] = []
    for i, slide in enumerate(prs.slides, 1):
        out.append(f"[Slayt {i}]")
        for sh in _walk_shapes(slide.shapes):
            if sh.has_text_frame:
                t = sh.text_frame.text.strip()
                if t:
                    out.append(t)
            if getattr(sh, "has_table", False) and sh.has_table:
                for row in sh.table.rows:
                    cells = _dedupe_cells([c.text.strip().replace("\n", " ") for c in row.cells])
                    if any(cells):
                        out.append(" | ".join(cells))
        if slide.has_notes_slide:
            n = slide.notes_slide.notes_text_frame.text.strip()
            if n:
                out.append("Konuşmacı notu: " + n)
    return Extracted("\n".join(out), page_count=len(prs.slides))


def _normalize(text: str) -> str:
    text = text.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t\f\v]+\n", "\n", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def extract(kind: str, data: bytes) -> Extracted:
    """Senkron ve CPU yoğun: çağıran thread pool içinde çalıştırmalıdır."""
    s = get_settings()
    res = {"pdf": _pdf, "docx": _docx, "pptx": _pptx}[kind](data)
    res.text = _normalize(res.text)
    if len(res.text) < 5:
        if kind == "pdf":
            if get_ocr_fn() is None:
                raise ExtractError("Bu PDF taranmış (görüntü) bir belge gibi görünüyor ve OCR kullanılamadığı için metin okunamadı.")
            raise ExtractError("Taranmış belgeden okunaklı bir metin çıkarılamadı.")
        raise ExtractError("Bu belgeden metin çıkarılamadı (içerik boş olabilir).")
    if len(res.text) > s.max_text_chars:
        res.text = res.text[: s.max_text_chars]
        res.truncated = True
        extra = f"Belge çok uzun olduğu için ilk {s.max_text_chars:,} karakter alındı.".replace(",", ".")
        res.note = f"{res.note} {extra}" if res.note else extra
    return res
